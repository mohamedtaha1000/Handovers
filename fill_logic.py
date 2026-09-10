#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
fill_logic.py
=============
The document-filling engine for every handover template. Started as a
single hardcoded function for the laptop-handover document; now holds one
fill function per template plus a TEMPLATES registry that the web app
uses to build the "pick a document" screen and each one's own form.

Every fill function edits only the specific data runs inside its
template's XML (never the fixed Arabic/English labels around them), so
the original formatting, fonts, RTL layout and table structure are
preserved exactly - only the data changes. It also auto-detects whether
each value is Arabic or Latin script and fixes that run's right-to-left
flag to match (this is what fixes the original "/" and Arabic-label-
gluing bugs - see fill_handover_script.md in the project for the story).
"""

import re
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE_DIR = Path(__file__).resolve().parent
DOC_TEMPLATES_DIR = BASE_DIR / "doc_templates"

WEEKDAYS_AR = ["الاثنين", "الثلاثاء", "الاربعاء", "الخميس", "الجمعة", "السبت", "الاحد"]
MONTHS_AR = [
    "يناير", "فبراير", "مارس", "ابريل", "مايو", "يونيو",
    "يوليو", "اغسطس", "سبتمبر", "اكتوبر", "نوفمبر", "ديسمبر",
]

ARABIC_RE = re.compile(r"[؀-ۿ]")

DEFAULT_COMPANY = "اس تي ام للاستثمار"


# ----------------------------------------------------------------------
# Low-level run helpers (shared by every fill function)
# ----------------------------------------------------------------------

def has_arabic(text):
    return bool(ARABIC_RE.search(text))


def set_value(run, text):
    """Set a run's text, and fix its right-to-left flag to match the
    script of the new text (Arabic vs Latin/digits).

    Defensive: a handful of source runs in the original templates had no
    <w:rPr> at all, which meant they silently inherited the document's
    docDefaults font size (12pt) instead of the 9pt used everywhere else
    in these tables - that mismatch is what caused values like a long
    job title to wrap awkwardly instead of fitting the column. If a run
    is missing rPr, give it one at the same 9pt size the rest of the
    table uses instead of leaving it to chance."""
    run.text = text
    rPr = run._r.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        run._r.insert(0, rPr)
        for tag, attrib in (
            ("w:rFonts", {"w:ascii": "Calibri", "w:hAnsi": "Calibri", "w:cs": "Calibri"}),
            ("w:sz", {"w:val": "18"}),
            ("w:szCs", {"w:val": "18"}),
        ):
            el = OxmlElement(tag)
            for k, v in attrib.items():
                el.set(qn(k), v)
            rPr.append(el)
    rtl_el = rPr.find(qn("w:rtl"))
    if has_arabic(text):
        if rtl_el is None:
            rPr.append(OxmlElement("w:rtl"))
    else:
        if rtl_el is not None:
            rPr.remove(rtl_el)


def clear(run):
    run.text = ""


def strip_hyperlinks(paragraph):
    p = paragraph._p
    for hl in p.findall(qn("w:hyperlink")):
        p.remove(hl)


def runs_of(table, row, col, para=0):
    return table.rows[row].cells[col].paragraphs[para].runs


def date_parts(date_obj):
    date_str = f"{date_obj.day}/{date_obj.month}/{date_obj.year}"
    day_name = WEEKDAYS_AR[date_obj.weekday()]
    month_name = MONTHS_AR[date_obj.month - 1]
    return date_str, day_name, month_name


# ----------------------------------------------------------------------
# Laptop handover
# ----------------------------------------------------------------------

def fill_laptop_handover(template_path, output_path, data):
    d = docx.Document(str(template_path))
    date_str, day_name, month_name = date_parts(data["date_obj"])

    name, department, role = data["name"], data["department"], data["role"]
    mobile, email, code, govid = data["mobile"], data["email"], data["code"], data["govid"]
    serial, model, cpu = data["serial"], data["model"], data["cpu"]
    company, color, storage, ram = data["company"], data["color"], data["storage"], data["ram"]

    t1, t2, t3, t4, t5 = d.tables[1], d.tables[2], d.tables[3], d.tables[4], d.tables[5]

    r = runs_of(t1, 0, 0)
    set_value(r[0], date_str)
    clear(r[1]); clear(r[2]); clear(r[3])
    r = runs_of(t1, 0, 1)
    set_value(r[3], month_name)
    r = runs_of(t1, 0, 2)
    set_value(r[1], day_name)
    clear(r[2])

    r = runs_of(t2, 0, 0)
    set_value(r[0], role + " ")
    r = runs_of(t2, 0, 1)
    set_value(r[2], name)
    clear(r[3])
    r = runs_of(t2, 1, 0)
    set_value(r[2], " " + company)
    r = runs_of(t2, 1, 1)
    set_value(r[2], department + " ")
    p = t2.rows[2].cells[0].paragraphs[0]
    r = p.runs
    set_value(r[0], mobile)
    clear(r[1])
    set_value(r[4], email)
    strip_hyperlinks(p)
    r = runs_of(t2, 2, 1)
    set_value(r[0], code)
    clear(r[1])

    set_value(runs_of(t3, 1, 0)[0], model)
    set_value(runs_of(t3, 1, 1)[0], serial)
    set_value(runs_of(t3, 1, 2)[0], color)
    set_value(runs_of(t3, 1, 3)[0], storage)
    set_value(runs_of(t3, 1, 4)[0], ram)
    r = runs_of(t3, 1, 5)
    set_value(r[0], cpu)
    if len(r) > 1:
        clear(r[1])

    r = runs_of(t4, 0, 0)
    set_value(r[3], govid)
    clear(r[4])
    set_value(r[13], company)
    clear(r[14])
    r = runs_of(t4, 0, 1)
    set_value(r[2], name)
    clear(r[3])
    set_value(r[26], department + " ")

    r = runs_of(t5, 2, 0)
    set_value(r[0], date_str)
    clear(r[1]); clear(r[2]); clear(r[3])
    r = runs_of(t5, 2, 1)
    set_value(r[0], date_str)
    clear(r[1]); clear(r[2]); clear(r[3])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(output_path))
    return output_path


# ----------------------------------------------------------------------
# Laptop replacement
# ----------------------------------------------------------------------

def fill_laptop_replacement(template_path, output_path, data):
    d = docx.Document(str(template_path))
    date_str, day_name, month_name = date_parts(data["date_obj"])

    name, department, role = data["name"], data["department"], data["role"]
    mobile, email, code, govid = data["mobile"], data["email"], data["code"], data["govid"]
    company = data["company"]

    t0, t1, t2, t3, t4, t5, t6, t7 = d.tables

    r = runs_of(t1, 0, 0)
    set_value(r[2], date_str)
    clear(r[3]); clear(r[4]); clear(r[5]); clear(r[6])
    r = runs_of(t1, 0, 1)
    set_value(r[3], month_name)
    r = runs_of(t1, 0, 2)
    set_value(r[2], day_name)

    r = runs_of(t2, 0, 0)
    set_value(r[2], role + " ")
    r = runs_of(t2, 0, 1)
    set_value(r[3], name)
    r = runs_of(t2, 1, 0)
    set_value(r[2], company)
    clear(r[3])
    r = runs_of(t2, 1, 1)
    set_value(r[0], department + " ")
    p = t2.rows[2].cells[0].paragraphs[0]
    r = p.runs
    set_value(r[1], mobile)
    clear(r[0])
    p2 = t2.rows[2].cells[0].paragraphs[1]
    r2 = p2.runs
    set_value(r2[1], email)
    clear(r2[0])
    strip_hyperlinks(p2)
    r = runs_of(t2, 2, 1)
    set_value(r[1], code)

    # NOTE: in the source document, table 3 ("Tablet Model" header, the
    # simpler single-run cells) sits under the "جهاز اللاب توب الجديد"
    # (NEW device) heading, and table 4 ("Laptop Model" header, the
    # messier multi-run cells) sits under "جهاز اللاب توب القديم" (OLD
    # device) - the table headers are misleading but the section
    # headings actually printed above each table are what matters, so
    # table 3 gets the new device's data and table 4 gets the old
    # device's - each keeping its own run layout below.
    r = runs_of(t3, 1, 0); set_value(r[1], data["new_model"])
    r = runs_of(t3, 1, 1); set_value(r[0], data["new_serial"])
    r = runs_of(t3, 1, 2); set_value(r[0], data["new_color"])
    r = runs_of(t3, 1, 3); set_value(r[0], data["new_storage"])
    r = runs_of(t3, 1, 4); set_value(r[0], data["new_ram"])
    r = runs_of(t3, 1, 5); set_value(r[0], data["new_cpu"])

    r = runs_of(t4, 1, 0)
    set_value(r[1], data["old_model"])
    clear(r[2]); clear(r[3]); clear(r[4])
    r = runs_of(t4, 1, 1); set_value(r[0], data["old_serial"]); clear(r[1])
    r = runs_of(t4, 1, 2); set_value(r[0], data["old_color"])
    r = runs_of(t4, 1, 3); set_value(r[0], data["old_storage"])
    r = runs_of(t4, 1, 4)
    set_value(r[0], data["old_ram"]); clear(r[1]); clear(r[2])
    r = runs_of(t4, 1, 5); set_value(r[0], data["old_cpu"]); clear(r[1])

    r = runs_of(t5, 0, 0)
    set_value(r[2], govid)
    clear(r[3])
    r = runs_of(t5, 0, 1)
    set_value(r[2], " " + name)
    r = runs_of(t5, 1, 0)
    set_value(r[0], company + " ")
    r = runs_of(t5, 1, 1)
    set_value(r[0], department + " ")

    r = runs_of(t7, 2, 0)
    set_value(r[2], date_str)
    clear(r[3])
    r = runs_of(t7, 2, 1)
    set_value(r[2], date_str)
    clear(r[3])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(output_path))
    return output_path


# ----------------------------------------------------------------------
# Keyboard / mouse (small peripheral) receipts - same document family,
# but not quite identical run layout underneath, so each gets its own
# explicit fill function (safer than guessing a shared run index).
# ----------------------------------------------------------------------

def fill_keyboard_receipt(template_path, output_path, data):
    d = docx.Document(str(template_path))
    date_str, day_name, month_name = date_parts(data["date_obj"])

    name, department, role = data["name"], data["department"], data["role"]
    mobile, code, govid = data["mobile"], data["code"], data["govid"]
    company, model, brand = data["company"], data["model"], data["brand"]

    t0, t1, t2, t3, t4, t5 = d.tables

    r = runs_of(t1, 0, 0); set_value(r[0], date_str)
    r = runs_of(t1, 0, 1); set_value(r[3], month_name)
    r = runs_of(t1, 0, 2); set_value(r[1], day_name)

    r = runs_of(t2, 0, 0); set_value(r[4], role)
    r = runs_of(t2, 0, 1); set_value(r[1], " " + name)
    r = runs_of(t2, 1, 0); set_value(r[0], "لدى شركة/ " + company + " ")
    r = runs_of(t2, 1, 1); set_value(r[2], department)
    r = runs_of(t2, 2, 0); set_value(r[3], mobile)
    r = runs_of(t2, 2, 1); set_value(r[0], code)

    r = runs_of(t3, 1, 3); set_value(r[0], model)
    r = runs_of(t3, 1, 4); set_value(r[0], brand)

    r = runs_of(t4, 0, 0); set_value(r[1], govid)
    r = runs_of(t4, 0, 1); set_value(r[4], name)
    r = runs_of(t4, 1, 0); set_value(r[0], "لدى شركة/ " + company + " ")
    r = runs_of(t4, 1, 1); set_value(r[1], department)

    r = runs_of(t5, 2, 0); set_value(r[0], date_str)
    r = runs_of(t5, 2, 1); set_value(r[0], date_str)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(output_path))
    return output_path


def fill_mouse_receipt(template_path, output_path, data):
    d = docx.Document(str(template_path))
    date_str, day_name, month_name = date_parts(data["date_obj"])

    name, department, role = data["name"], data["department"], data["role"]
    mobile, code, govid = data["mobile"], data["code"], data["govid"]
    company, model, brand = data["company"], data["model"], data["brand"]

    t0, t1, t2, t3, t4, t5 = d.tables

    r = runs_of(t1, 0, 0); set_value(r[0], date_str)
    r = runs_of(t1, 0, 1); set_value(r[3], month_name)
    r = runs_of(t1, 0, 2); set_value(r[3], day_name)

    r = runs_of(t2, 0, 0); set_value(r[4], role)
    r = runs_of(t2, 0, 1); set_value(r[1], " " + name)
    r = runs_of(t2, 1, 0); set_value(r[0], "لدى شركة/ " + company + " ")
    r = runs_of(t2, 1, 1); set_value(r[2], department)
    r = runs_of(t2, 2, 0); set_value(r[4], mobile)
    r = runs_of(t2, 2, 1); set_value(r[1], code)

    r = runs_of(t3, 1, 3); set_value(r[0], model)
    r = runs_of(t3, 1, 4); set_value(r[0], brand)

    r = runs_of(t4, 0, 0); set_value(r[1], govid)
    r = runs_of(t4, 0, 1); set_value(r[1], " " + name)
    r = runs_of(t4, 1, 0); set_value(r[0], "لدى شركة/ " + company + " ")
    r = runs_of(t4, 1, 1); set_value(r[3], department)

    r = runs_of(t5, 2, 0); set_value(r[0], date_str)
    r = runs_of(t5, 2, 1); set_value(r[0], date_str)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(output_path))
    return output_path


# ----------------------------------------------------------------------
# Screen / monitor handover
# ----------------------------------------------------------------------

def fill_screen_handover(template_path, output_path, data):
    d = docx.Document(str(template_path))
    date_str, day_name, month_name = date_parts(data["date_obj"])

    name, department, role = data["name"], data["department"], data["role"]
    mobile, code, govid = data["mobile"], data["code"], data["govid"]
    company = data["company"]
    model, serial, color = data["model"], data["serial"], data["color"]

    t0, t1, t2, t3, t4, t5 = d.tables

    r = runs_of(t1, 0, 0)
    set_value(r[1], date_str)
    clear(r[2]); clear(r[3]); clear(r[4])
    r = runs_of(t1, 0, 1)
    set_value(r[2], month_name)
    r = runs_of(t1, 0, 2)
    set_value(r[1], day_name)

    r = runs_of(t2, 0, 0)
    set_value(r[3], role)
    r = runs_of(t2, 0, 1)
    set_value(r[3], name)
    r = runs_of(t2, 1, 0)
    set_value(r[0], "لدى شركة/ " + company + " ")
    r = runs_of(t2, 1, 1)
    set_value(r[1], department)
    r = runs_of(t2, 2, 0)
    set_value(r[5], mobile)
    r = runs_of(t2, 2, 1)
    set_value(r[1], code)

    r = runs_of(t3, 1, 0)
    set_value(r[0], model)
    clear(r[1]); clear(r[2]); clear(r[3])
    r = runs_of(t3, 1, 2)
    set_value(r[1], serial)
    r = runs_of(t3, 1, 5)
    set_value(r[0], color)

    r = runs_of(t4, 0, 0)
    set_value(r[3], govid)
    r = runs_of(t4, 0, 1)
    set_value(r[3], name)
    r = runs_of(t4, 1, 0)
    set_value(r[0], "لدى شركة/ " + company + " ")
    r = runs_of(t4, 1, 1)
    set_value(r[7], department)

    r = runs_of(t5, 2, 0)
    set_value(r[2], date_str)
    clear(r[3]); clear(r[4]); clear(r[5]); clear(r[6])
    r = runs_of(t5, 2, 1)
    set_value(r[2], date_str)
    clear(r[3]); clear(r[4]); clear(r[5])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(output_path))
    return output_path


# ----------------------------------------------------------------------
# Template registry - the single source of truth for the "pick a
# document" screen and each one's own form fields.
# ----------------------------------------------------------------------

# Field specs shared by (almost) every template.
_NAME = {"key": "name", "label": "Full name", "placeholder": "e.g. Ahmed Ali Mohamed", "required": True}
_DEPARTMENT = {"key": "department", "label": "Department", "placeholder": "IT, Finance, Sales…", "required": True}
_ROLE = {"key": "role", "label": "Position", "placeholder": "Software Engineer", "required": True}
_MOBILE = {"key": "mobile", "label": "Mobile", "placeholder": "01xxxxxxxxx", "required": True, "inputmode": "numeric"}
_EMAIL = {"key": "email", "label": "Email", "required": True, "suffix": "@stm.com.eg"}
_CODE = {"key": "code", "label": "Employee code", "required": True}
_GOVID = {"key": "govid", "label": "National ID", "required": True, "maxlength": 14, "inputmode": "numeric", "placeholder": "14 digits", "wide": True}

EMPLOYEE_FIELDS_WITH_EMAIL = [_NAME, _DEPARTMENT, _ROLE, _MOBILE, _EMAIL, _CODE, _GOVID]
EMPLOYEE_FIELDS_NO_EMAIL = [_NAME, _DEPARTMENT, _ROLE, _MOBILE, _CODE, _GOVID]

TEMPLATES = {
    "laptop_handover": {
        "label": "Laptop handover",
        "label_ar": "محضر تسليم لاب توب",
        "doc_file": "Laptop Handover Template.docx",
        "filename_prefix": "استلام لابتوب",
        "fill": fill_laptop_handover,
        "employee_fields": EMPLOYEE_FIELDS_WITH_EMAIL,
        "device_title": "Laptop details",
        "device_fields": [
            {"key": "model", "label": "Laptop model", "placeholder": "E14", "required": True},
            {"key": "serial", "label": "Serial number", "placeholder": "PF5K...", "required": True},
            {"key": "cpu", "label": "CPU", "placeholder": "I5 / I7", "required": True},
        ],
        "extra_fields": [
            {"key": "color", "label": "Color", "placeholder": "BLACK", "default": "BLACK"},
            {"key": "storage", "label": "Storage", "placeholder": "512SSD", "default": "512SSD"},
            {"key": "ram", "label": "RAM", "placeholder": "24GB", "default": "24GB"},
            {"key": "company", "label": "Company", "placeholder": DEFAULT_COMPANY, "default": DEFAULT_COMPANY},
        ],
    },
    "laptop_replacement": {
        "label": "Laptop replacement",
        "label_ar": "محضر استبدال جهاز كمبيوتر",
        "doc_file": "Laptop Replacement Template.docx",
        "filename_prefix": "استبدال لابتوب",
        "fill": fill_laptop_replacement,
        "employee_fields": EMPLOYEE_FIELDS_WITH_EMAIL,
        "device_title": "Old device (being returned)",
        "device_fields": [
            {"key": "old_model", "label": "Old model", "placeholder": "E14", "required": True},
            {"key": "old_serial", "label": "Old serial number", "placeholder": "PF5K...", "required": True},
            {"key": "old_color", "label": "Old color", "placeholder": "BLACK", "required": True},
            {"key": "old_storage", "label": "Old storage", "placeholder": "512SSD", "required": True},
            {"key": "old_ram", "label": "Old RAM", "placeholder": "24GB", "required": True},
            {"key": "old_cpu", "label": "Old CPU", "placeholder": "I5 / I7", "required": True},
        ],
        "device_title_2": "New device (being issued)",
        "device_fields_2": [
            {"key": "new_model", "label": "New model", "placeholder": "E14", "required": True},
            {"key": "new_serial", "label": "New serial number", "placeholder": "PF5K...", "required": True},
            {"key": "new_color", "label": "New color", "placeholder": "BLACK", "required": True},
            {"key": "new_storage", "label": "New storage", "placeholder": "512SSD", "required": True},
            {"key": "new_ram", "label": "New RAM", "placeholder": "24GB", "required": True},
            {"key": "new_cpu", "label": "New CPU", "placeholder": "I5 / I7", "required": True},
        ],
        "extra_fields": [
            {"key": "company", "label": "Company", "placeholder": DEFAULT_COMPANY, "default": DEFAULT_COMPANY},
        ],
    },
    "keyboard_receipt": {
        "label": "Keyboard receipt",
        "label_ar": "محضر استلام كيبورد",
        "doc_file": "Keyboard Receipt Template.docx",
        "filename_prefix": "استلام كيبورد",
        "fill": fill_keyboard_receipt,
        "employee_fields": EMPLOYEE_FIELDS_NO_EMAIL,
        "device_title": "Keyboard details",
        "device_fields": [
            {"key": "model", "label": "Model", "placeholder": "9320M", "required": True},
            {"key": "brand", "label": "Brand", "placeholder": "Rapoo", "required": True},
        ],
        "extra_fields": [
            {"key": "company", "label": "Company", "placeholder": DEFAULT_COMPANY, "default": DEFAULT_COMPANY},
        ],
    },
    "mouse_receipt": {
        "label": "Mouse receipt",
        "label_ar": "محضر استلام ماوس",
        "doc_file": "Mouse Receipt Template.docx",
        "filename_prefix": "استلام ماوس",
        "fill": fill_mouse_receipt,
        "employee_fields": EMPLOYEE_FIELDS_NO_EMAIL,
        "device_title": "Mouse details",
        "device_fields": [
            {"key": "model", "label": "Model", "placeholder": "M171", "required": True},
            {"key": "brand", "label": "Brand", "placeholder": "Logitech", "required": True},
        ],
        "extra_fields": [
            {"key": "company", "label": "Company", "placeholder": DEFAULT_COMPANY, "default": DEFAULT_COMPANY},
        ],
    },
    "screen_handover": {
        "label": "Screen handover",
        "label_ar": "محضر تسليم شاشة",
        "doc_file": "Screen Handover Template.docx",
        "filename_prefix": "تسليم شاشة",
        "fill": fill_screen_handover,
        "employee_fields": EMPLOYEE_FIELDS_NO_EMAIL,
        "device_title": "Screen details",
        "device_fields": [
            {"key": "model", "label": "Screen model", "placeholder": "SE2725HMc (Dell 27 Monitor)", "required": True},
            {"key": "serial", "label": "Serial number", "placeholder": "CN-...", "required": True},
            {"key": "color", "label": "Color", "placeholder": "Black", "required": True},
        ],
        "extra_fields": [
            {"key": "company", "label": "Company", "placeholder": DEFAULT_COMPANY, "default": DEFAULT_COMPANY},
        ],
    },
}


def all_fields(template_id):
    """Every field key this template's form should collect (employee +
    device + device_2 if any), in order."""
    spec = TEMPLATES[template_id]
    fields = list(spec["employee_fields"]) + list(spec["device_fields"])
    if "device_fields_2" in spec:
        fields += list(spec["device_fields_2"])
    return fields


def required_field_keys(template_id):
    return [f["key"] for f in all_fields(template_id) if f.get("required")]


def template_path(template_id):
    return DOC_TEMPLATES_DIR / TEMPLATES[template_id]["doc_file"]
